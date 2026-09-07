from io import BytesIO
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook
from app.core.config import settings
from app.db.session import SessionLocal
from app.models.entities import ExportJob
from app.services import exports as export_service


def test_template_permissions_and_field_validation(client,engineer,admin):
    templates=client.get("/api/v1/export-templates",headers=engineer)
    assert templates.status_code==200
    assert any(item["is_default"] for item in templates.json()["data"])
    denied=client.post("/api/v1/export-templates",headers=engineer,json={"name":"越权模板","fields":["title"]})
    assert denied.status_code==403
    invalid=client.post("/api/v1/export-templates",headers=admin,json={"name":"无效字段模板","fields":["password_hash"]})
    assert invalid.status_code==422


def test_async_excel_export_and_download(client,engineer,manager):
    templates=client.get("/api/v1/export-templates",headers=engineer).json()["data"]
    template_id=next(item["id"] for item in templates if item["is_default"])
    created_record=client.post("/api/v1/records",headers=engineer,json={"content":"=HYPERLINK(\"https://invalid.example\",\"不应执行\")","title":"公式注入测试"})
    record_id=created_record.json()["data"]["id"]
    created=client.post("/api/v1/exports",headers=engineer,json={"template_id":template_id,"record_ids":[record_id]})
    assert created.status_code==202
    job_id=created.json()["data"]["id"]
    detail=client.get(f"/api/v1/exports/{job_id}",headers=engineer)
    assert detail.status_code==200
    assert detail.json()["data"]["status"]=="COMPLETED"
    assert detail.json()["data"]["record_count"]==1
    assert client.get(f"/api/v1/exports/{job_id}",headers=manager).status_code==404
    download=client.get(f"/api/v1/exports/{job_id}/download",headers=engineer)
    assert download.status_code==200
    assert download.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook=load_workbook(BytesIO(download.content),data_only=False)
    sheet=workbook["记录汇总"]
    headers=[cell.value for cell in sheet[4]]
    assert headers[:3]==["记录标题","项目","创建人"]
    assert sheet.freeze_panes=="A5"
    assert sheet.auto_filter.ref
    assert all(cell.data_type!="f" for row in sheet.iter_rows(min_row=5) for cell in row)
    notifications=client.get("/api/v1/notifications",headers=engineer).json()["data"]
    assert any(item["type"]=="EXPORT_COMPLETED" for item in notifications)


def test_export_failure_auto_retry_and_manual_retry(client,engineer,monkeypatch):
    template_id=client.get("/api/v1/export-templates",headers=engineer).json()["data"][0]["id"]
    original=export_service.save_export
    monkeypatch.setattr(settings,"export_job_max_attempts",2)
    monkeypatch.setattr(export_service,"save_export",lambda *_:(_ for _ in ()).throw(OSError("storage unavailable")))
    created=client.post("/api/v1/exports",headers=engineer,json={"template_id":template_id})
    job_id=created.json()["data"]["id"]
    failed=client.get(f"/api/v1/exports/{job_id}",headers=engineer).json()["data"]
    assert failed["status"]=="FAILED"
    assert failed["attempt_count"]==2
    monkeypatch.setattr(export_service,"save_export",original)
    retried=client.post(f"/api/v1/exports/{job_id}/retry",headers=engineer)
    assert retried.status_code==202
    completed=client.get(f"/api/v1/exports/{job_id}",headers=engineer).json()["data"]
    assert completed["status"]=="COMPLETED"
    assert completed["attempt_count"]==1


def test_expired_export_cleanup_and_regeneration(client,engineer,admin):
    template_id=client.get("/api/v1/export-templates",headers=engineer).json()["data"][0]["id"]
    created=client.post("/api/v1/exports",headers=engineer,json={"template_id":template_id})
    job_id=created.json()["data"]["id"]
    with SessionLocal() as db:
        job=db.get(ExportJob,job_id); job.expires_at=datetime.now(timezone.utc)-timedelta(minutes=1); db.commit()
    denied=client.post("/api/v1/exports/maintenance/cleanup",headers=engineer)
    assert denied.status_code==403
    cleaned=client.post("/api/v1/exports/maintenance/cleanup",headers=admin)
    assert cleaned.status_code==200
    assert cleaned.json()["data"]["cleaned"]>=1
    expired=client.get(f"/api/v1/exports/{job_id}",headers=engineer).json()["data"]
    assert expired["status"]=="EXPIRED"
    assert client.get(f"/api/v1/exports/{job_id}/download",headers=engineer).status_code==410
    assert client.post(f"/api/v1/exports/{job_id}/retry",headers=engineer).status_code==202
    assert client.get(f"/api/v1/exports/{job_id}",headers=engineer).json()["data"]["status"]=="COMPLETED"


def test_runtime_status_is_safe_and_admin_only(client,engineer,admin):
    assert client.get("/api/v1/system/runtime").status_code==401
    assert client.get("/api/v1/system/runtime",headers=engineer).status_code==403
    response=client.get("/api/v1/system/runtime",headers=admin)
    assert response.status_code==200
    data=response.json()["data"]
    assert data["queue"]["mode"]=="inline"
    assert data["queue"]["worker_online"] is True
    assert "redis_url" not in str(data)
