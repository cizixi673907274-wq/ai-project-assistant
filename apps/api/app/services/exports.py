from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from minio import Minio
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.core.config import settings
from app.models.entities import ExportJob, Notification, Record

PRIORITY_LABELS={"LOW":"低","MEDIUM":"中","HIGH":"高","URGENT":"紧急"}
STATUS_LABELS={"DRAFT":"草稿","SUBMITTED":"已提交","AI_PROCESSING":"AI分析中","AI_REVIEW_REQUIRED":"AI待确认","READY_TO_ASSIGN":"待分派","ASSIGNED":"已分派","IN_PROGRESS":"处理中","WAITING_SUPPLEMENT":"待补充","RESOLVED":"待确认解决","CLOSED":"已关闭","REJECTED":"已驳回","CANCELLED":"已撤销"}

FIELD_DEFINITIONS={
    "title":("记录标题",lambda r:r.title),
    "content":("原始记录",lambda r:r.content),
    "summary":("AI 摘要",lambda r:r.summary or ""),
    "project_name":("项目",lambda r:r.project.name if r.project else "未关联项目"),
    "creator_name":("创建人",lambda r:r.creator.name),
    "category":("问题分类",lambda r:r.category_snapshot or ""),
    "stage":("发生环节",lambda r:r.stage_snapshot or ""),
    "product":("涉及产品",lambda r:r.product_snapshot or ""),
    "priority":("优先级",lambda r:PRIORITY_LABELS.get(r.priority.value,r.priority.value)),
    "status":("处理状态",lambda r:STATUS_LABELS.get(r.status.value,r.status.value)),
    "created_at":("创建时间",lambda r:r.created_at),
    "updated_at":("更新时间",lambda r:r.updated_at),
}
DEFAULT_FIELDS=["title","project_name","creator_name","category","product","priority","status","summary","created_at"]


def validate_fields(fields:list[str])->list[str]:
    clean=[]
    for field in fields:
        if field not in FIELD_DEFINITIONS: raise ValueError(f"不支持的导出字段：{field}")
        if field not in clean: clean.append(field)
    if not clean: raise ValueError("导出模板至少需要一个字段")
    return clean


def _safe_cell(value):
    if isinstance(value,str) and value.startswith(("=","+","-","@")): return "'"+value
    if isinstance(value,datetime):
        if not value.tzinfo: value=value.replace(tzinfo=timezone.utc)
        value=value.astimezone(ZoneInfo("Asia/Shanghai")).replace(tzinfo=None)
        return value
    return value


def build_records_workbook(records:list[Record],fields:list[str],template_name:str,generated_by:str)->BytesIO:
    fields=validate_fields(fields); workbook=Workbook(); sheet=workbook.active; sheet.title="记录汇总"; sheet.sheet_view.showGridLines=False
    end=get_column_letter(len(fields)); now=datetime.now(timezone.utc)
    sheet.merge_cells(f"A1:{end}1"); sheet["A1"]="AI项目助手 · 研发记录导出"; sheet["A1"].font=Font(name="Microsoft YaHei",size=18,bold=True,color="FFFFFF"); sheet["A1"].fill=PatternFill("solid",fgColor="169B7A"); sheet["A1"].alignment=Alignment(vertical="center"); sheet.row_dimensions[1].height=36
    sheet.merge_cells(f"A2:{end}2"); sheet["A2"]=f"模板：{template_name}    记录数：{len(records)}    生成者：{generated_by}    生成时间：{now.astimezone().strftime('%Y-%m-%d %H:%M')}"; sheet["A2"].font=Font(name="Microsoft YaHei",size=10,color="5F706A"); sheet["A2"].alignment=Alignment(vertical="center"); sheet.row_dimensions[2].height=24
    headers=[FIELD_DEFINITIONS[field][0] for field in fields]
    for col,label in enumerate(headers,1):
        cell=sheet.cell(4,col,label); cell.font=Font(name="Microsoft YaHei",bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F8F73"); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    sheet.row_dimensions[4].height=28
    thin=Side(style="thin",color="DCE7E3")
    for row_index,record in enumerate(records,5):
        for col_index,field in enumerate(fields,1):
            value=_safe_cell(FIELD_DEFINITIONS[field][1](record)); cell=sheet.cell(row_index,col_index,value); cell.font=Font(name="Microsoft YaHei",size=10,color="1F2D2A"); cell.alignment=Alignment(vertical="top",wrap_text=field in {"content","summary","title","status"}); cell.border=Border(bottom=thin)
            if field in {"created_at","updated_at"}: cell.number_format="yyyy-mm-dd hh:mm"
        sheet.row_dimensions[row_index].height=38 if any(field in fields for field in ("content","summary")) else 24
    widths={"title":28,"content":48,"summary":48,"project_name":24,"creator_name":12,"category":18,"stage":16,"product":18,"priority":10,"status":14,"created_at":19,"updated_at":19}
    for index,field in enumerate(fields,1): sheet.column_dimensions[get_column_letter(index)].width=widths.get(field,16)
    sheet.freeze_panes="A5"; sheet.auto_filter.ref=f"A4:{end}{max(4,len(records)+4)}"; sheet.print_title_rows="1:4"; sheet.sheet_properties.pageSetUpPr.fitToPage=True; sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=0; sheet.sheet_properties.outlinePr.summaryBelow=True
    output=BytesIO(); workbook.save(output); output.seek(0); return output


def _minio(): return Minio(settings.minio_endpoint,access_key=settings.minio_access_key,secret_key=settings.minio_secret_key,secure=settings.minio_secure)


def save_export(file_key:str,data:bytes):
    if settings.export_storage=="minio":
        storage=_minio()
        if not storage.bucket_exists(settings.minio_bucket): storage.make_bucket(settings.minio_bucket)
        storage.put_object(settings.minio_bucket,file_key,BytesIO(data),len(data),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return
    target=(Path(settings.export_dir).resolve()/file_key).resolve(); root=Path(settings.export_dir).resolve()
    if root not in target.parents: raise ValueError("非法导出文件路径")
    target.parent.mkdir(parents=True,exist_ok=True); target.write_bytes(data)


def read_export(file_key:str)->bytes:
    if settings.export_storage=="minio":
        response=_minio().get_object(settings.minio_bucket,file_key)
        try: return response.read()
        finally: response.close(); response.release_conn()
    target=(Path(settings.export_dir).resolve()/file_key).resolve(); root=Path(settings.export_dir).resolve()
    if root not in target.parents or not target.is_file(): raise FileNotFoundError(file_key)
    return target.read_bytes()


def delete_export(file_key:str):
    if settings.export_storage=="minio":
        _minio().remove_object(settings.minio_bucket,file_key); return
    target=(Path(settings.export_dir).resolve()/file_key).resolve(); root=Path(settings.export_dir).resolve()
    if root not in target.parents: raise ValueError("非法导出文件路径")
    if target.is_file(): target.unlink()


def run_export_job(job_id:str,session_factory):
    db=session_factory()
    try:
        job=db.get(ExportJob,job_id)
        if not job or job.status!="PENDING": return job.status if job else "MISSING"
        job.status="PROCESSING"; job.attempt_count+=1; job.error=None; db.commit()
        record_ids=(job.filters_json or {}).get("record_ids",[])
        records=list(db.scalars(select(Record).where(Record.id.in_(record_ids),Record.deleted_at.is_(None)).order_by(Record.created_at.desc())).unique().all()) if record_ids else []
        content=build_records_workbook(records,job.template.fields_json,job.template.name,job.creator.name).getvalue(); filename=f"研发记录汇总-{datetime.now():%Y%m%d-%H%M%S}.xlsx"; file_key=f"exports/{job.creator_id}/{job.id}/{filename}"
        save_export(file_key,content); job.status="COMPLETED"; job.file_key=file_key; job.filename=filename; job.record_count=len(records); job.completed_at=datetime.now(timezone.utc); job.expires_at=job.completed_at+timedelta(days=settings.export_retention_days); db.add(Notification(user_id=job.creator_id,type="EXPORT_COMPLETED",title="Excel 导出已完成",content=f"《{filename}》已生成，可前往模板中心下载。")); db.commit(); return "COMPLETED"
    except Exception as exc:
        db.rollback(); job=db.get(ExportJob,job_id)
        if job:
            job.error=str(exc)[:4000]
            if job.attempt_count<settings.export_job_max_attempts: job.status="PENDING"; db.commit(); return "RETRY"
            job.status="FAILED"; job.completed_at=datetime.now(timezone.utc); db.add(Notification(user_id=job.creator_id,type="EXPORT_FAILED",title="Excel 导出失败",content="导出任务多次重试仍未成功，可在模板中心手动重试。")); db.commit()
        return "FAILED"
    finally: db.close()


def cleanup_expired_exports(session_factory,now:datetime|None=None)->int:
    db=session_factory(); cleaned=0; current=now or datetime.now(timezone.utc)
    try:
        jobs=list(db.scalars(select(ExportJob).where(ExportJob.status=="COMPLETED")).all())
        for job in jobs:
            if not job.expires_at:
                base=job.completed_at or job.created_at
                if not base.tzinfo: base=base.replace(tzinfo=timezone.utc)
                job.expires_at=base+timedelta(days=settings.export_retention_days)
            expires=job.expires_at if job.expires_at.tzinfo else job.expires_at.replace(tzinfo=timezone.utc)
            if expires>current: continue
            if job.file_key:
                try: delete_export(job.file_key)
                except FileNotFoundError: pass
            job.status="EXPIRED"; job.file_key=None; cleaned+=1
        db.commit(); return cleaned
    finally: db.close()
