from __future__ import annotations

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.models.entities import Project, Record, RecordEvent, RecordTask


GREEN = "169B7A"
DARK = "1F2D2A"
MUTED = "66756F"
LIGHT = "EFF8F5"
GRID = "DDE7E3"
STATUS_NAMES = {
    "DRAFT": "草稿",
    "SUBMITTED": "已提交",
    "AI_PROCESSING": "AI分析中",
    "AI_REVIEW_REQUIRED": "AI待确认",
    "WAITING_SUPPLEMENT": "待补充",
    "READY_TO_ASSIGN": "待分派",
    "ASSIGNED": "已分派",
    "IN_PROGRESS": "处理中",
    "RESOLVED": "待确认解决",
    "CLOSED": "已关闭",
    "REJECTED": "已驳回",
    "CANCELLED": "已取消",
}
PRIORITY_NAMES = {"LOW": "低", "MEDIUM": "中", "HIGH": "高", "URGENT": "紧急"}
EVENT_NAMES = {
    "SEEDED": "种子记录建立",
    "RECORD_CREATED": "创建记录",
    "RECORD_SUBMITTED": "提交记录",
    "AI_PROCESSING": "AI开始分析",
    "AI_ANALYSIS_COMPLETED": "AI分析完成",
    "AI_CONFIRMED": "确认AI结果",
    "AI_RESULT_EDITED": "修改AI结果",
    "RECORD_ASSIGNED": "分派责任人",
    "PROCESS_STARTED": "开始处理",
    "PROCESS_COMMENTED": "提交处理进展",
    "RECORD_RESOLVED": "标记解决",
    "RECORD_CLOSED": "关闭记录",
    "CLOSE_NOTE": "填写关闭说明",
}


def _font(run, size=11, bold=False, color=DARK):
    run.font.name = "Arial Unicode MS"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def _cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            tc_mar.append(node)


def _set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    if tbl_ind.getparent() is None:
        tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _cell_margins(cell)


def _write_cell(cell, value, bold=False, color=DARK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    _font(p.add_run(str(value or "-")), 10, bold, color)


def _heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), 16, True, GREEN)
    return p


def _body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    _font(p.add_run(text or "-"), 11)
    return p


def _enum_value(value):
    return getattr(value, "value", value)


def _brief(value: str | None, limit: int = 180):
    text = (value or "-").strip()
    return text if len(text) <= limit else f"{text[:limit].rstrip()}…"


def _configure_business_document(doc: Document):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance, section.footer_distance = Inches(0.45), Inches(0.45)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    for key in ("eastAsia", "ascii", "hAnsi"):
        normal._element.rPr.rFonts.set(qn(f"w:{key}"), "Arial Unicode MS")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6)):
        style = doc.styles[style_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size, style.font.bold = Pt(size), True
        style.font.color.rgb = RGBColor.from_string(GREEN)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
    return section


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(paragraph.add_run("AI项目助手 · 研发问题闭环报告　｜　第 "), 9, False, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    _font(paragraph.add_run(" 页"), 9, False, MUTED)


def build_project_report(project: Project, records: list[Record], generated_by: str, generated_at: datetime | None = None) -> BytesIO:
    """Generate the standard_business_brief project closure report."""
    generated_at = generated_at or datetime.now()
    doc = Document()
    section = _configure_business_document(doc)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(header.add_run(f"AI项目助手　｜　{project.code}"), 9, False, MUTED)
    _add_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    _font(kicker.add_run("PROJECT CLOSURE BRIEF"), 10, True, GREEN)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _font(title.add_run(f"{project.name} 研发问题闭环报告"), 24, True, DARK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    _font(subtitle.add_run(f"项目编号：{project.code}　｜　状态：{'进行中' if project.status == 'ACTIVE' else '已归档'}　｜　生成：{generated_at:%Y-%m-%d %H:%M} / {generated_by}"), 10, False, MUTED)

    closed = sum(_enum_value(r.status) in {"CLOSED", "CANCELLED"} for r in records)
    high = sum(_enum_value(r.priority) in {"HIGH", "URGENT"} for r in records)
    open_count = len(records) - closed
    completion = round(closed / len(records) * 100) if records else 100
    metrics = doc.add_table(rows=2, cols=4)
    metrics.style = "Table Grid"
    for cell, value in zip(metrics.rows[0].cells, ("记录总数", "待闭环", "高风险", "闭环率")):
        _shade(cell, LIGHT); _write_cell(cell, value, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER)
    for cell, value in zip(metrics.rows[1].cells, (len(records), open_count, high, f"{completion}%")):
        _write_cell(cell, value, True, DARK, WD_ALIGN_PARAGRAPH.CENTER)
    _set_table_geometry(metrics, [2340, 2340, 2340, 2340])

    _heading(doc, "1. 管理摘要")
    summary = f"本报告覆盖研发项目当前 {len(records)} 条产品开发记录，其中 {open_count} 条仍待闭环，{high} 条为高或紧急优先级，整体闭环率为 {completion}%。"
    summary += " 建议优先处理结构干涉、软件联调、模具试模等高优先级事项，并在关闭前补齐责任人处理说明。" if open_count else " 当前记录已全部闭环，建议归档关键经验并沉淀为研发知识库。"
    _body(doc, summary)

    _heading(doc, "2. 问题闭环清单")
    if records:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("序号", "记录 / 摘要", "分类", "优先级", "状态", "责任人 / 截止")):
            _shade(cell, GREEN); _write_cell(cell, value, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, record in enumerate(records, 1):
            active = next((task for task in getattr(record, "report_tasks", []) if _enum_value(task.status) not in {"DONE", "CANCELLED"}), None)
            values = (index, f"{record.title}\n{_brief(record.summary or record.content)}", record.category_snapshot or "待识别", PRIORITY_NAMES.get(str(_enum_value(record.priority)), str(_enum_value(record.priority))), STATUS_NAMES.get(str(_enum_value(record.status)), str(_enum_value(record.status))), f"{active.assignee.name if active else '未分派'}\n{active.due_at.strftime('%Y-%m-%d') if active and active.due_at else '未设截止'}")
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for column, (cell, value) in enumerate(zip(row.cells, values)):
                _write_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if column in {0,3,4} else WD_ALIGN_PARAGRAPH.LEFT)
                if index % 2 == 0: _shade(cell, "F7FAF9")
        _set_table_geometry(table, [520, 3100, 1350, 850, 1100, 2440])
    else:
        _body(doc, "当前项目暂无记录。")

    _heading(doc, "3. 风险与下一步")
    urgent = [r for r in records if _enum_value(r.priority) in {"HIGH", "URGENT"} and _enum_value(r.status) not in {"CLOSED", "CANCELLED"}]
    if urgent:
        for index, record in enumerate(urgent[:8], 1):
            _body(doc, f"{index}. {record.title}：{_brief(record.summary or record.content, 220)}")
    else:
        _body(doc, "未发现尚未闭环的高风险记录。建议持续跟踪待处理事项的责任人和截止时间。")
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    note.paragraph_format.space_after = Pt(0)
    _font(note.add_run("报告说明　"), 9, True, GREEN)
    _font(note.add_run("本报告由 AI项目助手基于系统实时记录自动生成，用于项目沟通与阶段归档；记录状态、责任人与处理结论以系统最新数据为准。"), 9, False, MUTED)

    output = BytesIO(); doc.save(output); output.seek(0); return output


def build_record_document(record: Record, events: list[RecordEvent], actor_names: dict[str, str], task: RecordTask | None = None) -> BytesIO:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 12, 8, 4)):
        style = doc.styles[style_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(GREEN if style_name != "Heading 3" else DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(header.add_run("AI项目助手 · 项目记录"), 9, False, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run("本文件由 AI项目助手 自动生成"), 9, False, MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    _font(kicker.add_run("PROJECT RECORD"), 10, True, GREEN)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    _font(title.add_run("研发记录归档"), 24, True, DARK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    _font(subtitle.add_run(record.title), 14, True, MUTED)

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta_rows = [
        ("记录编号", record.id),
        ("所属项目", record.project.name if record.project else "未关联项目"),
        ("提交人", record.creator.name),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for row, (label, value) in zip(meta.rows, meta_rows):
        _shade(row.cells[0], LIGHT)
        _write_cell(row.cells[0], label, True, GREEN)
        _write_cell(row.cells[1], value)
    _set_table_geometry(meta, [2100, 7260])

    _heading(doc, "1. 记录概览")
    overview = doc.add_table(rows=3, cols=2)
    overview.style = "Table Grid"
    status = STATUS_NAMES.get(str(_enum_value(record.status)), str(_enum_value(record.status)))
    priority = PRIORITY_NAMES.get(str(_enum_value(record.priority)), str(_enum_value(record.priority)))
    overview_rows = [("当前状态", status), ("优先级", priority), ("创建时间", record.created_at.strftime("%Y-%m-%d %H:%M"))]
    for row, (label, value) in zip(overview.rows, overview_rows):
        _shade(row.cells[0], LIGHT)
        _write_cell(row.cells[0], label, True, GREEN)
        _write_cell(row.cells[1], value)
    _set_table_geometry(overview, [2100, 7260])

    _heading(doc, "2. 原始记录")
    _body(doc, record.content)

    _heading(doc, "3. AI 结构化分析")
    analysis = doc.add_table(rows=5, cols=2)
    analysis.style = "Table Grid"
    confidence = f"{round(record.analyses[-1].confidence * 100)}%" if record.analyses else "待分析"
    analysis_rows = [
        ("AI 摘要", record.summary or "待分析"),
        ("问题分类", record.category_snapshot or "待识别"),
        ("发生环节", record.stage_snapshot or "待识别"),
        ("涉及产品", record.product_snapshot or "待识别"),
        ("AI 置信度", confidence),
    ]
    for row, (label, value) in zip(analysis.rows, analysis_rows):
        _shade(row.cells[0], LIGHT)
        _write_cell(row.cells[0], label, True, GREEN)
        _write_cell(row.cells[1], value)
    _set_table_geometry(analysis, [2100, 7260])

    _heading(doc, "4. 处理结果")
    result = doc.add_table(rows=2, cols=2)
    result.style = "Table Grid"
    result_rows = [("当前责任人", task.assignee.name if task else "未分派"), ("处理说明", task.response if task and task.response else "暂无处理说明")]
    for row, (label, value) in zip(result.rows, result_rows):
        _shade(row.cells[0], LIGHT)
        _write_cell(row.cells[0], label, True, GREEN)
        _write_cell(row.cells[1], value)
    _set_table_geometry(result, [2100, 7260])

    _heading(doc, "5. 附件与媒体")
    if record.media:
        media_table = doc.add_table(rows=1, cols=3)
        media_table.style = "Table Grid"
        for cell, value in zip(media_table.rows[0].cells, ("文件名", "类型", "大小")):
            _shade(cell, GREEN)
            _write_cell(cell, value, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        for media in record.media:
            cells = media_table.add_row().cells
            _write_cell(cells[0], media.original_name)
            _write_cell(cells[1], _enum_value(media.type), align=WD_ALIGN_PARAGRAPH.CENTER)
            _write_cell(cells[2], f"{media.size / 1024:.1f} KB", align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_table_geometry(media_table, [5200, 2200, 1960])
    else:
        _body(doc, "暂无附件。")

    _heading(doc, "6. 处理流程")
    flow = doc.add_table(rows=1, cols=4)
    flow.style = "Table Grid"
    for cell, value in zip(flow.rows[0].cells, ("时间", "节点", "操作人", "说明")):
        _shade(cell, GREEN)
        _write_cell(cell, value, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    for item in events:
        cells = flow.add_row().cells
        detail = item.after_json or item.before_json or {}
        _write_cell(cells[0], item.created_at.strftime("%Y-%m-%d\n%H:%M"), align=WD_ALIGN_PARAGRAPH.CENTER)
        _write_cell(cells[1], EVENT_NAMES.get(item.event_type, item.event_type))
        _write_cell(cells[2], actor_names.get(item.actor_id or "", "系统"), align=WD_ALIGN_PARAGRAPH.CENTER)
        _write_cell(cells[3], "；".join(f"{key}: {value}" for key, value in detail.items()) or "-")
    flow.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    _set_table_geometry(flow, [1800, 2100, 1560, 3900])

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def build_records_table_document(records: list[Record]) -> BytesIO:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    section.header_distance = section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(header.add_run("AI项目助手 · 记录汇总"), 9, False, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run("本文件由 AI项目助手 自动生成"), 9, False, MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    _font(kicker.add_run("PROJECT RECORD REGISTER"), 9, True, GREEN)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _font(title.add_run("研发记录批量汇总表"), 22, True, DARK)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    _font(meta.add_run(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　｜　记录数量：{len(records)} 条"), 10, False, MUTED)

    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ("序号", "记录标题", "所属项目", "提交人", "问题分类", "优先级", "状态", "记录内容 / AI摘要")
    for cell, value in zip(table.rows[0].cells, headers):
        _shade(cell, GREEN)
        _write_cell(cell, value, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    for index, record in enumerate(records, 1):
        cells = table.add_row().cells
        values = (
            index,
            record.title,
            record.project.name if record.project else "未关联项目",
            record.creator.name,
            record.category_snapshot or "待识别",
            PRIORITY_NAMES.get(str(_enum_value(record.priority)), str(_enum_value(record.priority))),
            STATUS_NAMES.get(str(_enum_value(record.status)), str(_enum_value(record.status))),
            record.summary or record.content,
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _write_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if column in {0,3,5,6} else WD_ALIGN_PARAGRAPH.LEFT)
            if index % 2 == 0:
                _shade(cell, "F7FAF9")
    _set_table_geometry(table, [600, 2500, 1750, 1000, 1600, 900, 1150, 4468])

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
